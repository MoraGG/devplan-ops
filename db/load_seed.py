"""房地产开发运营计划系统 · 标准库数据加载器（参数化，可读 Excel 直插 MySQL）。
作为 02_seed.sql 的可执行替代：避开字符串转义，直接 executemany 写入。
"""
import os
import openpyxl, re, pymysql

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "template", "远建2026版三级计划（绿城基础含公式）.xlsx")
wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb["节点计划 (终稿)"]
rows = list(ws.iter_rows(min_row=4, values_only=True))

LEVEL_MAP = {'里程碑': '里程碑', '一级': '一级', '二级': '二级', '三级': '二级*'}
BT = [('BT18', '18F高层'), ('BT4', '4F及以下'), ('BT6', '6F叠拼'), ('BT8', '8F洋房'),
      ('BT11', '11F小高层'), ('BT27', '27F高层'), ('BT33', '33F高层')]
PRE = [('volum', '容积率', '["小于2.5","大于2.5"]'), ('basefloor', '地下室层数', '[1,2,3]'),
       ('pile', '桩基形式', '["管桩","灌注桩"]'), ('hsupp', '水平支撑层数', '[0,1]'),
       ('demoarea', '首开示范区面积', '["小于7000㎡","大于7000㎡"]'), ('voidlayer', '有无架空层', '["有","无"]'),
       ('topfloor', '地上最高层数', '["数值"]'), ('devloan', '有无开发贷', '["有","无"]'),
       ('delivery', '交付形式', '["毛坯","精装"]'), ('interior', '有无户内改造', '["有","无"]'),
       ('facade', '外立面形式', '["保温涂料","保温涂料+铝板"]')]

def parse_formula(f):
    """把 L 列公式解析为「依赖行 + 偏移天数」。

    支持三类形态（取公式中【首个】L 引用，与 extract_rules 的分支提取口径一致）：
      =L6+30 / =L82-150          带偏移引用
      =L53                       纯引用 -> 偏移 0
      =IF(前置信息!$B$10="精装",L139-150,"")  取 L139 -150，条件分支交由 std_node_rule
    ⚠️ 历史缺陷：旧正则 r'L(\\d+)([+\\-])(\\d+)' 强制要求带 ±数字，
      导致「正式开工（开工）」的 =L53 被判为无依赖（depend_row=None）而丢边。
    """
    if not isinstance(f, str) or not f.startswith('='):
        return None
    m = re.search(r'L(\d+)(?:([+\-])(\d+))?', f)
    if m:
        row = int(m.group(1))
        off = int(m.group(3)) * (1 if m.group(2) == '+' else -1) if m.group(2) else 0
        return dict(depend_row=row, offset_days=off, base_is_t0=int(row == 6))
    return dict(depend_row=None, offset_days=None, base_is_t0=0)

# 解析
profs = []; profset = set()
for r in rows:
    if r[3] in LEVEL_MAP and r[5]:
        p = str(r[5]).strip()
        if p not in profset:
            profset.add(p); profs.append(p)
profmap = {p: i + 1 for i, p in enumerate(profs)}

nodes = []; rownum_map = {}
for idx, r in enumerate(rows):
    if r[3] not in LEVEL_MAP:
        continue
    er = idx + 4
    level = LEVEL_MAP[r[3]]
    lvl_seq = r[0] if r[3] == '里程碑' else (r[1] if r[3] == '一级' else (r[2] if r[3] in ('二级', '三级') else None))
    fm = parse_formula(r[11])
    nodes.append(dict(excel_row=er, level=level, seq=lvl_seq, name=str(r[4]).strip() if r[4] else None,
                      prof=str(r[5]).strip() if r[5] else None, deliverable=r[19], support_doc=r[20],
                      duration=r[7], formula=r[11], **(fm or {})))
    rownum_map[er] = len(nodes)

DB_NAME = os.environ.get("DB_NAME", "dev_plan")
# autocommit=True：否则 information_schema 查询会开长事务并持有 MDL，
# 一旦进程中断，后续 DROP/CREATE TABLE 会被 "Waiting for table metadata lock" 阻塞。
con = pymysql.connect(host=os.environ.get("DB_HOST", "127.0.0.1"),
                      port=int(os.environ.get("DB_PORT", "3306")),
                      user=os.environ.get("DB_USER", "root"),
                      password=os.environ.get("DB_PASS", ""),
                      database=DB_NAME, charset="utf8mb4", autocommit=True)
cur = con.cursor()

# 清空
cur.execute("SET FOREIGN_KEY_CHECKS=0")
cur.execute("SELECT table_name FROM information_schema.tables WHERE table_schema=%s", (DB_NAME,))
for (t,) in cur.fetchall():
    cur.execute(f"DROP TABLE IF EXISTS {t}")
cur.execute("SET FOREIGN_KEY_CHECKS=1")
con.commit()

# 建表（剔除行首与【行内】注释后按 ; 拆分）
# ⚠️ 旧写法只删「以 -- 开头的行」，但 01_ddl.sql 的列注释里含 ';'
#    （如 "-- ...excel_row; NULL=该参数组合下节点不排程"），行内分号会把
#    CREATE TABLE 截断 → 1064 语法错误。故必须先按行内 -- 截断再拆分。
# 另：三个子表（project_node_rule / project_node_fixed / rule_change_log）
#    的外键指向后文才定义的 project，故建表期间关闭外键检查（已实测可豁免前向引用）。
ddl = open(os.path.join(HERE, "01_ddl.sql"), encoding="utf-8").read()
_lines = []
for ln in ddl.split("\n"):
    cut = ln.find("--")
    if cut >= 0:
        ln = ln[:cut]
    if ln.strip():
        _lines.append(ln)
_body = "\n".join(_lines)
stmts = [s.strip() for s in _body.split(";") if s.strip()]
print(f"DDL 语句数: {len(stmts)}")
cur.execute("SET FOREIGN_KEY_CHECKS=0")
for s in stmts:
    cur.execute(s)
cur.execute("SET FOREIGN_KEY_CHECKS=1")
con.commit()

# 参数化插入
cur.executemany("INSERT INTO profession(code,name) VALUES(%s,%s)",
                [(f"P{i+1}", p) for i, p in enumerate(profs)])
cur.executemany("INSERT INTO building_type(code,name) VALUES(%s,%s)", BT)
cur.executemany("INSERT INTO prereq_param(code,name,options_json) VALUES(%s,%s,%s)", PRE)

std_node_rows = []
for n in nodes:
    pid = profmap.get(n['prof']) if n['prof'] else None
    std_node_rows.append((n['excel_row'], n['level'], n['seq'], n['name'], pid,
                          n['deliverable'], n['support_doc'],
                          n['duration'] if isinstance(n['duration'], (int, float)) else None,
                          n['formula'], n.get('depend_row'), n.get('offset_days'),
                          n.get('base_is_t0', 0), n.get('has_condition', 0)))
cur.executemany(
    "INSERT INTO std_node(excel_row,level,seq,name,profession_id,deliverable,support_doc,duration,tmpl_finish_formula,depend_row,offset_days,base_is_t0,has_condition) "
    "VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", std_node_rows)

sd_rows = []
for i, n in enumerate(nodes):
    sid = i + 1
    for j, (c, _) in enumerate(BT):
        cell = rows[n['excel_row'] - 4][12 + j]
        if cell and str(cell).strip():
            sd_rows.append((sid, j + 1, str(cell).strip()))
cur.executemany("INSERT INTO std_duration(std_node_id,building_type_id,std_desc) VALUES(%s,%s,%s)", sd_rows)

dep_rows = []
for i, n in enumerate(nodes):
    if n.get('depend_row'):
        did = rownum_map.get(n['depend_row'])
        off = n['offset_days']; direction = '之后' if off >= 0 else '之前'
        dep_rows.append((i + 1, did, n['depend_row'], off, n.get('base_is_t0', 0), direction, n.get('has_condition', 0)))
cur.executemany(
    "INSERT INTO std_node_dependency(std_node_id,depend_std_node_id,depend_row,offset_days,base_is_t0,direction,has_condition) "
    "VALUES(%s,%s,%s,%s,%s,%s,%s)", dep_rows)

con.commit()
cur.execute("SELECT COUNT(*) FROM std_node"); print("std_node:", cur.fetchone()[0])
cur.execute("SELECT COUNT(*) FROM std_duration"); print("std_duration:", cur.fetchone()[0])
cur.execute("SELECT COUNT(*) FROM std_node_dependency"); print("std_node_dependency:", cur.fetchone()[0])
cur.execute("SELECT level,COUNT(*) FROM std_node GROUP BY level")
print("by level:", cur.fetchall())
cur.close(); con.close()
print("参数化加载完成。")
