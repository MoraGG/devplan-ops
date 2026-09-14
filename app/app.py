# -*- coding: utf-8 -*-
"""房地产开发运营计划系统 · Web 管理后台（MVP 正式前端，自包含版）

依赖：Flask + db/engine.py（相对工期引擎）+ db/dbconn.py（连接层，默认 SQLite）
运行：python3 app.py  ->  监听 $PORT（默认 5000），绑定 0.0.0.0
说明：发布/预览用默认 SQLite（DB_TYPE=sqlite，无需外部服务）；部署到自带 MySQL 的
      服务器时，设环境变量 DB_TYPE=mysql 即可切回生产库，代码无需改动。
"""
import os
import sys
import io
import datetime
import json
from flask import Flask, render_template, request, redirect, url_for, jsonify, send_file

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(__file__)), "db"))
from engine import RelativeDateEngine, ensure_default_alert_rules, compute_alerts  # noqa: E402
from dbconn import get_conn  # noqa: E402

TODAY = datetime.date(2026, 9, 11)  # 系统当前业务日期（演示用，可改为 datetime.date.today()）

app = Flask(__name__)
app.jinja_env.globals["TODAY"] = TODAY

# 启动即保证预警默认规则存在（里程碑/一级≥14天、二级/二级*≥7天，内控版）
ensure_default_alert_rules()


def db():
    return get_conn()


def engine_for(t0, prereq_json=None, project_id=None):
    return RelativeDateEngine(t0, prereq_json, project_id)


def parse_prereq_opts(options_json):
    """解析前置参数可选项：返回选项字符串列表；['数值'] 表示自由数值输入 -> 返回 None；其余 -> 空列表(文本)。"""
    try:
        opts = json.loads(options_json)
    except Exception:
        return []
    if isinstance(opts, list):
        if len(opts) == 1 and str(opts[0]) == "数值":
            return None  # 自由数值输入
        return [str(x) for x in opts]
    return []


def load_prereqs():
    """读取前置参数表，并解析出每个参数的下拉选项（opts）。"""
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,code,name,options_json FROM prereq_param ORDER BY id")
    prereqs = [dict(r) for r in cur.fetchall()]
    for pr in prereqs:
        pr["opts"] = parse_prereq_opts(pr["options_json"])
    cur.close(); con.close()
    return prereqs


def load_project(pid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name,client,building_type_id,plan_start,plan_deliver,status,prereq_json "
                "FROM project WHERE id=?", (pid,))
    p = cur.fetchone()
    cur.close(); con.close()
    return p


def building_types():
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name FROM building_type ORDER BY id")
    bts = cur.fetchall()
    cur.close(); con.close()
    return bts


# ---------------- 看板 / 项目列表 ----------------
@app.route("/")
def dashboard():
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name,client,building_type_id,plan_start,plan_deliver,status,prereq_json FROM project ORDER BY id")
    projects = cur.fetchall()
    cur.execute("SELECT id,name FROM building_type")
    bt = {r["id"]: r["name"] for r in cur.fetchall()}
    # 各项目预警计数（读取最近一次计算的结果）
    cur.execute("SELECT pn.project_id AS pid, COUNT(*) AS c FROM alert_record ar "
                "JOIN plan_node pn ON ar.plan_node_id=pn.id GROUP BY pn.project_id")
    alert_by_proj = {r["pid"]: r["c"] for r in cur.fetchall()}
    cur.close(); con.close()

    cards = []
    for p in projects:
        t0 = p["plan_start"]
        eng = engine_for(t0, p["prereq_json"], p["id"])
        d = eng.to_dict(project=p["name"])
        # 下一个未过期的里程碑
        upcoming = [m for m in d["milestones"] if m["date"] and m["date"] >= TODAY.isoformat()]
        next_ms = upcoming[0] if upcoming else None
        cards.append(dict(
            id=p["id"], name=p["name"], client=p["client"],
            building_type=bt.get(p["building_type_id"], "—"),
            t0=(t0.isoformat() if hasattr(t0, "isoformat") else str(t0)),
            delivery=d["delivery_date"],
            total_months=d["total_months"], total_days=d["total_days"],
            crit_nodes=len(d["critical_path"]),
            crit_miles=sum(1 for m in d["milestones"] if m["critical"]),
            status=p["status"],
            next_milestone=next_ms,
            alerts=alert_by_proj.get(p["id"], 0),
        ))
    kpi = dict(
        total=len(cards),
        on_going=sum(1 for c in cards if c["status"] == "进行中"),
        delivered=sum(1 for c in cards if c["status"] == "已交付"),
        avg_months=round(sum(c["total_months"] or 0 for c in cards) / len(cards), 1) if cards else 0,
        alerts=sum(alert_by_proj.values()),
    )
    return render_template("dashboard.html", cards=cards, kpi=kpi)


# ---------------- 新建项目 ----------------
@app.route("/project/new", methods=["GET", "POST"])
def new_project():
    bts = building_types()
    prereqs = load_prereqs()
    if request.method == "POST":
        con = db(); cur = con.cursor()
        name = request.form["name"].strip()
        client = request.form.get("client", "").strip()
        bt_id = int(request.form["building_type_id"])
        t0 = request.form["plan_start"]
        status = request.form.get("status", "进行中")
        buffer_kh = int(request.form.get("buffer_kh") or 0)
        buffer_ly = int(request.form.get("buffer_ly") or 0)
        prereq = {f"p{pr['id']}": request.form.get(f"p{pr['id']}") for pr in prereqs}
        prereq_json = json.dumps(prereq, ensure_ascii=False)
        # 交付日先按引擎预估（套模板后回填）
        eng = engine_for(t0, prereq_json)
        d = eng.to_dict(project=name)
        cur.execute(
            "INSERT INTO project(name,client,building_type_id,plan_start,plan_deliver,status,prereq_json,buffer_kh,buffer_ly) "
            "VALUES(?,?,?,?,?,?,?,?,?)",
            (name, client, bt_id, t0, d["delivery_date"], status, prereq_json, buffer_kh, buffer_ly))
        pid = cur.lastrowid
        con.commit(); cur.close(); con.close()
        # 生成三版（内控 + 考核 + 履约），考核/履约版 = 内控 + 项目级期量裕度
        generate_version(pid, "内控", 0)
        generate_version(pid, "考核", buffer_kh)
        generate_version(pid, "履约", buffer_ly)
        # 回写交付日
        con = db(); cur = con.cursor()
        cur.execute("UPDATE project SET plan_deliver=? WHERE id=?", (d["delivery_date"], pid))
        con.commit(); cur.close(); con.close()
        return redirect(url_for("project_detail", pid=pid))
    return render_template("new_project.html", bts=bts, prereqs=prereqs, existing={})


# ---------------- 编辑项目 ----------------
@app.route("/project/<int:pid>/edit", methods=["GET", "POST"])
def edit_project(pid):
    p = load_project(pid)
    if not p:
        return "项目不存在", 404
    bts = building_types()
    prereqs = load_prereqs()
    existing = json.loads(p["prereq_json"]) if p["prereq_json"] else {}
    if request.method == "POST":
        con = db(); cur = con.cursor()
        name = request.form["name"].strip()
        client = request.form.get("client", "").strip()
        bt_id = int(request.form["building_type_id"])
        t0 = request.form["plan_start"]
        status = request.form.get("status", "进行中")
        buffer_kh = int(request.form.get("buffer_kh") or 0)
        buffer_ly = int(request.form.get("buffer_ly") or 0)
        prereq = {f"p{pr['id']}": request.form.get(f"p{pr['id']}") for pr in prereqs}
        prereq_json = json.dumps(prereq, ensure_ascii=False)
        old_t0 = p["plan_start"]
        old_prereq = p["prereq_json"]
        old_kh = p.get("buffer_kh") or 0
        old_ly = p.get("buffer_ly") or 0
        cur.execute(
            "UPDATE project SET name=?,client=?,building_type_id=?,plan_start=?,status=?,prereq_json=?,buffer_kh=?,buffer_ly=? WHERE id=?",
            (name, client, bt_id, t0, status, prereq_json, buffer_kh, buffer_ly, pid))
        # 任一影响计划的因素改变 → 重新生成三版（保留 actual_*/status）
        if t0 != old_t0 or prereq_json != old_prereq or buffer_kh != old_kh or buffer_ly != old_ly:
            d = engine_for(t0, prereq_json, pid).to_dict(project=name)
            cur.execute("UPDATE project SET plan_deliver=? WHERE id=?", (d["delivery_date"], pid))
            con.commit(); cur.close(); con.close()
            regenerate_all_versions(pid)
        else:
            con.commit(); cur.close(); con.close()
            ensure_all_versions(pid)
        return redirect(url_for("project_detail", pid=pid))
    return render_template("edit_project.html", p=p, bts=bts, prereqs=prereqs, existing=existing)


# ---------------- 项目详情 ----------------
@app.route("/project/<int:pid>")
def project_detail(pid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name,client,building_type_id,plan_start,plan_deliver,status,prereq_json FROM project WHERE id=?", (pid,))
    p = cur.fetchone()
    if not p:
        cur.close(); con.close(); return "项目不存在", 404
    cur.execute("SELECT id,name FROM building_type")
    bt = {r["id"]: r["name"] for r in cur.fetchall()}
    # 懒生成缺失版本（保证三版齐备，保留已填报 actual_*）
    ensure_all_versions(pid)
    # 版本切换（?v=内控/考核/履约，默认内控）
    cur_v = request.args.get("v", "内控")
    if cur_v not in PLAN_TYPES:
        cur_v = "内控"
    vid = version_vid(pid, cur_v)
    # 计划节点（按 Excel 行序全量展示，随所选版本切换 finish_inner）
    cur.execute(
        "SELECT pn.id,pn.level,pn.seq,pn.name,pn.std_node_id,pn.finish_inner,pn.status,pr.name AS prof, sn.excel_row "
        "FROM plan_node pn LEFT JOIN profession pr ON pn.profession_id=pr.id "
        "LEFT JOIN std_node sn ON pn.std_node_id=sn.id "
        "WHERE pn.project_id=? AND pn.version_id=? ORDER BY sn.excel_row", (pid, vid))
    rows = cur.fetchall()
    cur.execute("SELECT DISTINCT std_node_id FROM std_node_rule")
    rule_ids = {r["std_node_id"] for r in cur.fetchall()}
    cur.execute("SELECT DISTINCT std_node_id FROM project_node_rule WHERE project_id=?", (pid,))
    rule_over = {r["std_node_id"] for r in cur.fetchall()}
    cur.execute("SELECT DISTINCT std_node_id FROM project_node_fixed WHERE project_id=?", (pid,))
    fixed_over = {r["std_node_id"] for r in cur.fetchall()}
    override_ids = rule_over | fixed_over
    cur.execute("SELECT id,excel_row,level,seq,name FROM std_node ORDER BY level,seq,excel_row")
    all_nodes = [dict(r) for r in cur.fetchall()]
    cur.execute("SELECT COUNT(*) AS c FROM alert_record ar JOIN plan_node pn ON ar.plan_node_id=pn.id "
                "WHERE pn.project_id=?", (pid,))
    alert_count = cur.fetchone()["c"]
    cur.close(); con.close()
    rule_params = [pr for pr in load_prereqs() if pr["code"] in ("topfloor", "devloan", "delivery", "interior", "facade")]

    # 引擎：关键路径（标准库图 + 本项目 T0 + 本项目前置参数 + 本项目 override）
    eng = engine_for(p["plan_start"], p["prereq_json"], p["id"])
    d = eng.to_dict(project=p["name"])
    crit_ids = {c["id"] for c in d["critical_path"]}

    # 里程碑距今天数（m.date 为 ISO 字符串或 date 对象）
    for m in d["milestones"]:
        md = m["date"]
        m["days_from_today"] = ((md - TODAY).days if isinstance(md, datetime.date)
                                else (datetime.date.fromisoformat(md) - TODAY).days) if md else None

    for r in rows:
        r["critical"] = r["std_node_id"] in crit_ids
        fi = r["finish_inner"]
        r["days_from_today"] = ((fi - TODAY).days if isinstance(fi, datetime.date)
                                else (datetime.date.fromisoformat(fi) - TODAY).days) if fi else None
    return render_template("project_detail.html", p=p, building_type=bt.get(p["building_type_id"], "—"),
                           data=d, rows=rows, rule_ids=rule_ids, override_ids=override_ids,
                           rule_params=rule_params, all_nodes=all_nodes, alert_count=alert_count,
                           cur_version=cur_v, plan_types=PLAN_TYPES,
                           crit_miles=sum(1 for m in d["milestones"] if m["critical"]))


@app.route("/project/<int:pid>/critical-path.json")
def critical_path_json(pid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name,plan_start,prereq_json FROM project WHERE id=?", (pid,))
    p = cur.fetchone(); cur.close(); con.close()
    if not p:
        return jsonify(error="not found"), 404
    eng = engine_for(p["plan_start"], p["prereq_json"], p["id"])
    return jsonify(eng.to_dict(project=p["name"]))


# ---------------- 标准节点库（可编辑，体现"规则数据化"） ----------------
@app.route("/std-nodes")
def std_nodes():
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,excel_row,level,seq,name,profession_id,tmpl_finish_formula,depend_row,offset_days,base_is_t0 FROM std_node ORDER BY excel_row")
    nodes = cur.fetchall()
    cur.execute("SELECT id,name FROM profession")
    prof = {r["id"]: r["name"] for r in cur.fetchall()}
    cur.execute("SELECT id,excel_row,level,seq,name FROM std_node")
    node_meta = {r["id"]: r for r in cur.fetchall()}
    cur.execute("SELECT std_node_id,depend_row,offset_days,base_is_t0,direction,depend_std_node_id FROM std_node_dependency")
    deps = {(r["std_node_id"]): r for r in cur.fetchall()}
    cur.execute("SELECT DISTINCT std_node_id FROM std_node_rule WHERE is_active=1")
    rule_ids = {r["std_node_id"] for r in cur.fetchall()}
    cur.execute("SELECT id,version_no,name,note,effective_from,operator,created_at FROM std_rule_version ORDER BY version_no DESC")
    versions = [dict(r) for r in cur.fetchall()]
    cur.close(); con.close()
    for n in nodes:
        n["prof"] = prof.get(n["profession_id"], "—")
        dp = deps.get(n["id"])
        if dp:
            if dp["base_is_t0"]:
                n["dep_desc"] = "T0基准" + (f" +{dp['offset_days']}天" if dp["offset_days"] else "")
            else:
                dn = node_meta.get(dp["depend_std_node_id"])
                dname = dn["name"] if dn else f"L{dp['depend_row']}"
                n["dep_desc"] = (f"依赖 {dname}(L{dp['depend_row']}) "
                                 f"{dp['direction']} {abs(dp['offset_days'])}天")
        else:
            n["dep_desc"] = "—"
        if n["id"] in rule_ids:
            n["dep_desc"] += " · 含参数化分支"
    return render_template("std_nodes.html", nodes=nodes, rule_ids=rule_ids, total=len(nodes), versions=versions)


@app.route("/std-node/<int:nid>/edit", methods=["GET", "POST"])
def edit_std_node(nid):
    con = db(); cur = con.cursor()
    if request.method == "POST":
        depend_row = request.form.get("depend_row") or None
        offset_days = int(request.form["offset_days"])
        direction = request.form.get("direction", "之后")
        base_is_t0 = 1 if request.form.get("base_is_t0") == "1" else 0
        # 主责专业：表单提交才更新，未提交（None/空）保持原值，避免误清空
        prof_raw = request.form.get("profession_id")
        profession_id = int(prof_raw) if prof_raw not in (None, "") else None
        off = offset_days if direction == "之后" else -offset_days
        # 解析 depend_row -> 真实依赖节点 id（引擎依赖 depend_std_node_id 做遍历）
        depend_std_node_id = None
        if depend_row and not base_is_t0:
            cur.execute("SELECT id FROM std_node WHERE excel_row=?", (int(depend_row),))
            r = cur.fetchone()
            depend_std_node_id = r["id"] if r else None
        # 工作副本语义：退役该节点旧生效行，插入新行（is_active=1, version_no=当前工作版）
        wv = cur_work_version(cur)
        cur.execute("UPDATE std_node_dependency SET is_active=0 WHERE std_node_id=? AND is_active=1", (nid,))
        cur.execute(
            "INSERT INTO std_node_dependency(std_node_id,depend_std_node_id,depend_row,offset_days,direction,base_is_t0,version_no,is_active) "
            "VALUES(?,?,?,?,?,?,?,1)", (nid, depend_std_node_id, depend_row, off, direction, base_is_t0, wv))
        # 同步 tmpl_finish_formula 文本（仅展示用）
        formula = "T0" if base_is_t0 else (f"=L{depend_row}{('+' if off>=0 else '')}{off}" if depend_row else None)
        if profession_id is not None:
            cur.execute("UPDATE std_node SET depend_row=?,offset_days=?,base_is_t0=?,tmpl_finish_formula=?,profession_id=? WHERE id=?",
                        (depend_row, off, base_is_t0, formula, profession_id, nid))
        else:
            cur.execute("UPDATE std_node SET depend_row=?,offset_days=?,base_is_t0=?,tmpl_finish_formula=? WHERE id=?",
                        (depend_row, off, base_is_t0, formula, nid))
        con.commit(); cur.close(); con.close()
        return redirect(url_for("std_nodes"))
    cur.execute("SELECT id,level,seq,name,profession_id,tmpl_finish_formula,depend_row,offset_days,base_is_t0 FROM std_node WHERE id=?", (nid,))
    node = cur.fetchone()
    cur.execute("SELECT std_node_id,depend_row,offset_days,base_is_t0,direction FROM std_node_dependency WHERE std_node_id=?", (nid,))
    dep = cur.fetchone()
    cur.execute("SELECT id,name FROM profession")
    profs = cur.fetchall()
    # 全部标准节点（供依赖下拉选择，按名称而非行号）
    cur.execute("SELECT id,excel_row,level,seq,name FROM std_node ORDER BY level,seq,excel_row")
    all_nodes = cur.fetchall()
    cur.execute("SELECT 1 FROM std_node_rule WHERE std_node_id=? LIMIT 1", (nid,))
    has_rule = bool(cur.fetchone())
    cur.close(); con.close()
    return render_template("edit_std_node.html", node=node, dep=dep, profs=profs, all_nodes=all_nodes, has_rule=has_rule)


# ---------------- 编辑参数化分支规则（IF/AND 条件，扁平分支卡片） ----------------
@app.route("/std-node/<int:nid>/rules", methods=["GET", "POST"])
def edit_std_node_rules(nid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,excel_row,level,seq,name FROM std_node WHERE id=?", (nid,))
    node = cur.fetchone()
    if not node:
        cur.close(); con.close(); return "节点不存在", 404
    cur.execute("SELECT id,excel_row,level,seq,name FROM std_node ORDER BY level,seq,excel_row")
    all_nodes = cur.fetchall()
    # 规则相关参数（仅 B8~B12 五个），带 opts
    prereqs = load_prereqs()
    rule_params = [pr for pr in prereqs if pr["code"] in ("topfloor", "devloan", "delivery", "interior", "facade")]
    # 现有分支
    cur.execute("""SELECT branch_order,is_default,conditions_json,depend_row,offset_expr,direction
                   FROM std_node_rule WHERE std_node_id=? ORDER BY branch_order""", (nid,))
    rules = []
    for r in cur.fetchall():
        try:
            conds = json.loads(r["conditions_json"]) if r["conditions_json"] else []
        except Exception:
            conds = []
        oe = r["offset_expr"] or ""
        if oe.lstrip("-").isdigit():
            val = int(oe)
            rules.append(dict(conditions=conds, depend_row=r["depend_row"],
                              offset_days=abs(val), direction=("之后" if val >= 0 else "之前"),
                              advanced="", is_default=bool(r["is_default"])))
        else:
            rules.append(dict(conditions=conds, depend_row=r["depend_row"],
                              offset_days=0, direction=r["direction"] or "之后",
                              advanced=oe, is_default=bool(r["is_default"])))
    cur.close(); con.close()

    if request.method == "POST":
        raw = request.form.get("rules_json", "[]")
        try:
            branches = json.loads(raw)
        except Exception:
            branches = []
        clean = []
        for b in branches:
            if not isinstance(b, dict):
                continue
            conds = []
            for cd in b.get("conditions", []):
                if not isinstance(cd, dict):
                    continue
                p = cd.get("param"); op = cd.get("op"); v = cd.get("value")
                if not p or op not in ("<=", ">=", ">", "<", "=", "<>"):
                    continue
                if v in (None, ""):
                    continue
                if p == "topfloor":
                    try:
                        v = int(v)
                    except Exception:
                        pass
                conds.append({"param": p, "op": op, "value": v})
            is_default = bool(b.get("is_default"))
            depend_row = b.get("depend_row") or None
            adv = (b.get("offset_expr") or "").strip()
            if adv:
                offset_expr = adv
            else:
                days = int(b.get("offset_days") or 0)
                direction = b.get("direction") or "之后"
                offset_expr = str(days) if direction == "之后" else str(-days)
            clean.append(dict(conditions=conds if not is_default else [], depend_row=depend_row,
                              offset_expr=offset_expr, direction=b.get("direction") or "之后",
                              is_default=is_default))
        before = rules_snapshot("标准库", nid)
        con = db(); cur = con.cursor()
        wv = cur_work_version(cur)
        cur.execute("UPDATE std_node_rule SET is_active=0 WHERE std_node_id=? AND is_active=1", (nid,))
        for i, b in enumerate(clean):
            cur.execute(
                "INSERT INTO std_node_rule(std_node_id,branch_order,is_default,conditions_json,depend_row,offset_expr,direction,base_is_t0,version_no,is_active) "
                "VALUES(?,?,?,?,?,?,?,0,?,1)",
                (nid, i + 1, 1 if b["is_default"] else 0,
                 json.dumps(b["conditions"], ensure_ascii=False), b["depend_row"], b["offset_expr"], b["direction"], wv))
        cur.execute("UPDATE std_node SET has_condition=1 WHERE id=?", (nid,))
        # 同步 std_node 的默认依赖展示（非参数化回退用）：取默认分支或首个分支
        default_b = next((b for b in clean if b["is_default"]), None) or (clean[0] if clean else None)
        if default_b:
            off = int(default_b["offset_expr"]) if default_b["offset_expr"].lstrip("-").isdigit() else 0
            cur.execute("UPDATE std_node SET depend_row=?,offset_days=?,base_is_t0=0 WHERE id=?",
                        (default_b["depend_row"], off, nid))
        con.commit(); cur.close(); con.close()
        log_rule_change("标准库", nid, None, "更新规则", before, clean_to_front(clean))
        return redirect(url_for("std_nodes"))

    con = db(); cur = con.cursor()
    cur.execute("SELECT id,action,operator,created_at FROM rule_change_log "
                "WHERE scope='标准库' AND std_node_id=? ORDER BY id DESC LIMIT 20", (nid,))
    logs = [dict(r) for r in cur.fetchall()]; cur.close(); con.close()
    return render_template("edit_std_node_rules.html", node=node, rules=rules,
                           all_nodes=all_nodes, rule_params=rule_params, logs=logs)


# ---------------- 项目级规则：分支清洗（与标准库共用） ----------------
def sanitize_branches(branches):
    clean = []
    for b in branches:
        if not isinstance(b, dict):
            continue
        conds = []
        for cd in b.get("conditions", []):
            if not isinstance(cd, dict):
                continue
            p = cd.get("param"); op = cd.get("op"); v = cd.get("value")
            if not p or op not in ("<=", ">=", ">", "<", "=", "<>"):
                continue
            if v in (None, ""):
                continue
            if p == "topfloor":
                try:
                    v = int(v)
                except Exception:
                    pass
            conds.append({"param": p, "op": op, "value": v})
        is_default = bool(b.get("is_default"))
        depend_row = b.get("depend_row") or None
        adv = (b.get("offset_expr") or "").strip()
        if adv:
            offset_expr = adv
        else:
            days = int(b.get("offset_days") or 0)
            direction = b.get("direction") or "之后"
            offset_expr = str(days) if direction == "之后" else str(-days)
        clean.append(dict(conditions=conds if not is_default else [], depend_row=depend_row,
                          offset_expr=offset_expr, direction=b.get("direction") or "之后",
                          is_default=is_default))
    return clean


def to_front(rs):
    """将 DB 行集（std_node_rule / project_node_rule）转前端分支列表。"""
    out = []
    for r in rs:
        conds = json.loads(r["conditions_json"]) if r["conditions_json"] else []
        oe = r["offset_expr"] or ""
        if oe.lstrip("-").isdigit():
            val = int(oe)
            out.append(dict(conditions=conds, depend_row=r["depend_row"], offset_days=abs(val),
                            direction=("之后" if val >= 0 else "之前"), advanced="", is_default=bool(r["is_default"])))
        else:
            out.append(dict(conditions=conds, depend_row=r["depend_row"], offset_days=0,
                            direction=r["direction"] or "之后", advanced=oe, is_default=bool(r["is_default"])))
    return out


def baseline_vid(pid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id FROM plan_version WHERE project_id=? AND plan_type='内控' ORDER BY id LIMIT 1", (pid,))
    r = cur.fetchone(); cur.close(); con.close()
    return r["id"] if r else None


def _as_date(v):
    """字符串/DATE → date 对象；无法解析返回 None。"""
    if v is None:
        return None
    if isinstance(v, datetime.date):
        return v
    try:
        return datetime.date.fromisoformat(str(v)[:10])
    except Exception:
        return None


def version_vid(pid, plan_type):
    """取某项目某 plan_type 的 plan_version id；无则 None（通用版 baseline_vid）。"""
    con = db(); cur = con.cursor()
    cur.execute("SELECT id FROM plan_version WHERE project_id=? AND plan_type=? ORDER BY id LIMIT 1",
                (pid, plan_type))
    r = cur.fetchone(); cur.close(); con.close()
    return r["id"] if r else None


def generate_version(pid, plan_type, buffer_days=0):
    """生成/重生成某版 plan_node（先删后插，整版覆盖）。基于当前生效引擎 + 项目 override。"""
    con = db(); cur = con.cursor()
    cur.execute("SELECT plan_start,prereq_json FROM project WHERE id=?", (pid,))
    pj = cur.fetchone(); con.close()
    if not pj:
        return None
    eng = engine_for(pj["plan_start"], pj["prereq_json"], pid)
    vid = version_vid(pid, plan_type)
    con = db(); cur = con.cursor()
    if vid is None:
        cur.execute("INSERT INTO plan_version(project_id,version_no,plan_type,status) VALUES(?,1,?,?)",
                    (pid, plan_type, '基准'))
        vid = cur.lastrowid
    else:
        # 先删本版本预警记录（alert_record.plan_node_id 外键约束 plan_node），否则 MySQL 下删 plan_node 报 1451
        cur.execute("DELETE FROM alert_record WHERE plan_node_id IN "
                    "(SELECT id FROM plan_node WHERE project_id=? AND version_id=?)", (pid, vid))
        cur.execute("DELETE FROM plan_node WHERE project_id=? AND version_id=?", (pid, vid))
    con.commit()
    eng.apply_to_project(pid, vid, plan_type, buffer_days=buffer_days)
    con.close()
    return vid


def regenerate_all_versions(pid):
    """内控/考核/履约三版统一重生成（保留已填报 actual_*/status），并刷新三版预警。"""
    con = db(); cur = con.cursor()
    cur.execute("SELECT buffer_kh,buffer_ly FROM project WHERE id=?", (pid,))
    b = cur.fetchone()
    # 备份现有 actual（按 std_node_id，三版值相同，取任意版即可）
    cur.execute("SELECT std_node_id,actual_start,actual_finish,progress_pct FROM plan_node "
                "WHERE project_id=?", (pid,))
    actual_map = {r["std_node_id"]: r for r in cur.fetchall()}
    con.close()
    kh = (b["buffer_kh"] or 0) if b else 0
    ly = (b["buffer_ly"] or 0) if b else 0
    for pt, buf in (("内控", 0), ("考核", kh), ("履约", ly)):
        vid = generate_version(pid, pt, buf)
        if vid is None or not actual_map:
            continue
        con = db(); cur = con.cursor()
        for snid, r in actual_map.items():
            cur.execute(
                "UPDATE plan_node SET actual_start=?,actual_finish=?,progress_pct=? "
                "WHERE project_id=? AND version_id=? AND std_node_id=?",
                (r["actual_start"], r["actual_finish"], r["progress_pct"], pid, vid, snid))
        con.commit(); con.close()
    for pt in ("内控", "考核", "履约"):
        compute_alerts(pid, pt, as_of=TODAY)


def ensure_all_versions(pid):
    """懒生成缺失的版本（仅建缺失的，不覆盖已存在版本，保留 actual_*/status）。"""
    con = db(); cur = con.cursor()
    cur.execute("SELECT buffer_kh,buffer_ly FROM project WHERE id=?", (pid,))
    b = cur.fetchone(); con.close()
    if not b:
        return
    kh = b["buffer_kh"] or 0
    ly = b["buffer_ly"] or 0
    if version_vid(pid, "内控") is None:
        generate_version(pid, "内控", 0)
    if version_vid(pid, "考核") is None:
        generate_version(pid, "考核", kh)
    if version_vid(pid, "履约") is None:
        generate_version(pid, "履约", ly)


def save_actual(pid, rows_actual):
    """同步写三版 plan_node 的 actual_*（同一事实，三版一致），并刷新三版预警。"""
    con = db(); cur = con.cursor()
    cur.execute("SELECT id FROM plan_version WHERE project_id=?", (pid,))
    vids = [r["id"] for r in cur.fetchall()]
    con.close()
    con = db(); cur = con.cursor()
    for vid in vids:
        for ra in rows_actual:
            cur.execute(
                "UPDATE plan_node SET actual_start=?,actual_finish=?,progress_pct=? "
                "WHERE project_id=? AND version_id=? AND std_node_id=?",
                (ra["actual_start"], ra["actual_finish"], ra["progress_pct"], pid, vid, ra["std_node_id"]))
    con.commit(); con.close()
    for pt in ("内控", "考核", "履约"):
        compute_alerts(pid, pt, as_of=TODAY)


def compare_rows(pid):
    """三版并排对比数据：节点元信息 + 内控/考核/履约 finish_inner + actual_finish + 偏差。"""
    con = db(); cur = con.cursor()
    vers = {}
    for pt in ("内控", "考核", "履约"):
        vid = version_vid(pid, pt)
        if vid is None:
            vers[pt] = {}
            continue
        cur.execute("SELECT std_node_id,finish_inner,actual_finish,status FROM plan_node "
                    "WHERE project_id=? AND version_id=?", (pid, vid))
        vers[pt] = {r["std_node_id"]: dict(r) for r in cur.fetchall()}
    cur.execute("SELECT sn.id,sn.excel_row,sn.level,sn.seq,sn.name,pr.name AS prof "
                "FROM std_node sn LEFT JOIN profession pr ON sn.profession_id=pr.id ORDER BY sn.excel_row")
    nodes = cur.fetchall()
    con.close()
    out = []
    for n in nodes:
        snid = n["id"]
        ne = vers["内控"].get(snid, {})
        kh = vers["考核"].get(snid, {})
        ly = vers["履约"].get(snid, {})
        fi_n = ne.get("finish_inner"); fi_k = kh.get("finish_inner"); fi_l = ly.get("finish_inner")
        af = ne.get("actual_finish")
        kh_diff = None; ly_diff = None; af_diff = None
        if fi_k and fi_n:
            kh_diff = (_as_date(fi_k) - _as_date(fi_n)).days
        if fi_l and fi_n:
            ly_diff = (_as_date(fi_l) - _as_date(fi_n)).days
        if af and fi_n:
            af_diff = (_as_date(af) - _as_date(fi_n)).days
        out.append(dict(
            std_node_id=snid, excel_row=n["excel_row"], level=n["level"], seq=n["seq"],
            name=n["name"], prof=n["prof"],
            fi_inner=fi_n, fi_kh=fi_k, fi_ly=fi_l, actual=af,
            kh_diff=kh_diff, ly_diff=ly_diff, af_diff=af_diff,
            status=ne.get("status")))
    return out


def build_compare_xlsx(pid):
    """生成项目三版计划对比工作簿（openpyxl），返回 BytesIO；项目不存在返回 None。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    ensure_all_versions(pid)
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name,client,plan_start,buffer_kh,buffer_ly FROM project WHERE id=?", (pid,))
    p = cur.fetchone()
    con.close()
    if not p:
        return None
    rows = compare_rows(pid)

    wb = Workbook(); ws = wb.active; ws.title = "三版计划对比"
    # 标题块
    ws.append([f"{p['name']}（{p['client'] or '—'}）· 三版计划对比"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    ws["A1"].font = Font(bold=True, size=13)
    sub = (f"项目启动：{p['plan_start']}　期量裕度：考核+{p['buffer_kh'] or 0}天 / "
           f"履约+{p['buffer_ly'] or 0}天　业务当前日：{TODAY}")
    ws.append([sub])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=11)
    ws["A2"].font = Font(size=10, color="666666")

    headers = ["序号", "等级", "节点名称", "内控完成日", "考核完成日", "履约完成日", "实际完成日",
               "内控→考核(天)", "内控→履约(天)", "实际−内控(天)", "状态"]
    ws.append(headers)
    hrow = 3
    hdr_fill = PatternFill("solid", fgColor="305496")
    for c in range(1, 12):
        cell = ws.cell(row=hrow, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in rows:
        fi_n = _as_date(r["fi_inner"]); fi_k = _as_date(r["fi_kh"])
        fi_l = _as_date(r["fi_ly"]); af = _as_date(r["actual"])
        status = r["status"] or "正常"
        ws.append([
            r["seq"] if r["seq"] is not None else "—",
            r["level"] or "", r["name"],
            fi_n, fi_k, fi_l, af,
            r["kh_diff"], r["ly_diff"], r["af_diff"], status,
        ])
        rr = ws.max_row
        for col in (4, 5, 6, 7):
            cell = ws.cell(row=rr, column=col)
            if isinstance(cell.value, datetime.date):
                cell.number_format = "YYYY-MM-DD"
        for col in (8, 9, 10):
            v = ws.cell(row=rr, column=col).value
            if isinstance(v, int):
                ws.cell(row=rr, column=col).value = (f"{v:+d}" if v != 0 else "一致")
        scell = ws.cell(row=rr, column=11)
        if status == "延期":
            scell.fill = PatternFill("solid", fgColor="F8CBAD")
        elif status == "预警":
            scell.fill = PatternFill("solid", fgColor="FFE699")
        for col in range(1, 12):
            ws.cell(row=rr, column=col).border = border
    widths = [8, 8, 40, 13, 13, 13, 13, 14, 14, 14, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A{hrow}:K{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def recompute_project_finish(pid, vid, t0, prereq_json):
    """按当前（含项目级 override 的）引擎重算该版本所有节点的 finish_inner（保留 actual 等填报），并回写交付日。"""
    if vid is None:
        return
    eng = engine_for(t0, prereq_json, pid)
    off = eng.compute_offsets()
    con = db(); cur = con.cursor()
    for nid, o in off.items():
        fin = (eng.t0 + datetime.timedelta(days=o)).isoformat() if o is not None else None
        cur.execute("UPDATE plan_node SET finish_inner=? WHERE project_id=? AND version_id=? AND std_node_id=?",
                    (fin, pid, vid, nid))
    d = eng.to_dict(project="")
    cur.execute("UPDATE project SET plan_deliver=? WHERE id=?", (d["delivery_date"], pid))
    con.commit(); cur.close(); con.close()


def log_rule_change(scope, std_node_id, project_id, action, before=None, after=None, operator="代建运营部", con=None):
    """记录标准库 / 项目级规则与固定日期的变更留痕。
    con 可传入调用方连接以复用事务（防 SQLite 写锁）；为 None 时自行开连接并提交。"""
    own = con is None
    if own:
        con = db()
    cur = con.cursor()
    cur.execute(
        "INSERT INTO rule_change_log(scope,std_node_id,project_id,action,before_json,after_json,operator) "
        "VALUES(?,?,?,?,?,?,?)",
        (scope, std_node_id, project_id, action,
         json.dumps(before, ensure_ascii=False) if before is not None else None,
         json.dumps(after, ensure_ascii=False) if after is not None else None, operator))
    if own:
        con.commit()
    cur.close()
    if own:
        con.close()


def rules_snapshot(scope, std_node_id, project_id=None):
    """读取某节点当前生效规则（前端友好格式），用于变更日志 before。"""
    con = db(); cur = con.cursor()
    if scope == "标准库":
        cur.execute("""SELECT branch_order,is_default,conditions_json,depend_row,offset_expr,direction
                       FROM std_node_rule WHERE std_node_id=? AND is_active=1 ORDER BY branch_order""", (std_node_id,))
    else:
        cur.execute("""SELECT branch_order,is_default,conditions_json,depend_row,offset_expr,direction
                       FROM project_node_rule WHERE project_id=? AND std_node_id=? ORDER BY branch_order""",
                    (project_id, std_node_id))
    rows = cur.fetchall(); cur.close(); con.close()
    return to_front(rows)


def clean_to_front(clean):
    """将 sanitize 后的 clean 分支列表转前端友好格式（供日志 after）。"""
    out = []
    for b in clean:
        oe = b.get("offset_expr") or ""
        if oe.lstrip("-").isdigit():
            val = int(oe)
            out.append(dict(conditions=b.get("conditions", []), depend_row=b.get("depend_row"),
                            offset_days=abs(val), direction="之后" if val >= 0 else "之前",
                            advanced="", is_default=bool(b.get("is_default"))))
        else:
            out.append(dict(conditions=b.get("conditions", []), depend_row=b.get("depend_row"),
                            offset_days=0, direction=b.get("direction") or "之后",
                            advanced=oe, is_default=bool(b.get("is_default"))))
    return out


# ---------------- 项目级编辑节点（弹窗交互）：继承 / 固定日期 / 修改规则 ----------------
@app.route("/project/<int:pid>/node/<int:nid>/edit", methods=["GET", "POST"])
def edit_project_node(pid, nid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,excel_row,level,seq,name FROM std_node WHERE id=?", (nid,))
    node = cur.fetchone()
    if not node:
        cur.close(); con.close(); return jsonify(error="node not found"), 404
    cur.execute("SELECT id,name,plan_start,prereq_json FROM project WHERE id=?", (pid,))
    proj = cur.fetchone()
    if not proj:
        cur.close(); con.close(); return jsonify(error="project not found"), 404
    cur.execute("SELECT fixed_date FROM project_node_fixed WHERE project_id=? AND std_node_id=?", (pid, nid))
    fr = cur.fetchone()
    cur.execute("SELECT branch_order,is_default,conditions_json,depend_row,offset_expr,direction "
                "FROM project_node_rule WHERE project_id=? AND std_node_id=? ORDER BY branch_order", (pid, nid))
    prr = cur.fetchall()
    cur.execute("SELECT 1 FROM std_node_rule WHERE std_node_id=? LIMIT 1", (nid,))
    has_rule = bool(cur.fetchone())
    cur.execute("SELECT finish_inner FROM plan_node WHERE project_id=? AND std_node_id=? ORDER BY id DESC LIMIT 1", (pid, nid))
    fin = cur.fetchone()
    con.close()

    if fr:
        mode = "fixed"; fixed_date = fr["fixed_date"]; rules = []
    elif prr:
        mode = "rule"; fixed_date = None; rules = to_front(prr)
    else:
        mode = "inherit"; fixed_date = None
        if has_rule:
            con = db(); cur = con.cursor()
            cur.execute("SELECT branch_order,is_default,conditions_json,depend_row,offset_expr,direction "
                        "FROM std_node_rule WHERE std_node_id=? ORDER BY branch_order", (nid,))
            rules = to_front(cur.fetchall()); con.close()
        else:
            # 普通节点（无参数化分支）：用标准库默认依赖构造单条默认分支，供"修改规则"初始化
            con = db(); cur = con.cursor()
            cur.execute("SELECT depend_row,offset_days FROM std_node WHERE id=?", (nid,))
            dn = cur.fetchone(); con.close()
            if dn and dn["depend_row"] is not None:
                off = int(dn["offset_days"] or 0)
                rules = [dict(conditions=[], depend_row=dn["depend_row"], offset_days=abs(off),
                              direction="之后" if off >= 0 else "之前", offset_expr="", is_default=True)]
            else:
                rules = [dict(conditions=[], depend_row=None, offset_days=0, direction="之后", offset_expr="", is_default=True)]

    if request.method == "POST":
        m = request.form.get("mode", "inherit")
        # 变更前快照（用于日志 before）
        if fr:
            before = {"fixed_date": fr["fixed_date"]}
        elif prr:
            before = to_front(prr)
        else:
            before = rules_snapshot("标准库", nid) if has_rule else {"depend_row": (rules[0]["depend_row"] if rules else None)}
        vid = baseline_vid(pid)
        con = db(); cur = con.cursor()
        if m == "inherit":
            cur.execute("DELETE FROM project_node_rule WHERE project_id=? AND std_node_id=?", (pid, nid))
            cur.execute("DELETE FROM project_node_fixed WHERE project_id=? AND std_node_id=?", (pid, nid))
            log_rule_change("项目", nid, pid, "恢复继承", before, None)
        elif m == "fixed":
            fixed_date = (request.form.get("fixed_date") or "").strip()
            try:
                datetime.date.fromisoformat(fixed_date)
            except Exception:
                con.close(); return jsonify(ok=False, msg="请填写合法的固定日期"), 400
            cur.execute("DELETE FROM project_node_rule WHERE project_id=? AND std_node_id=?", (pid, nid))
            cur.execute("DELETE FROM project_node_fixed WHERE project_id=? AND std_node_id=?", (pid, nid))
            cur.execute("INSERT INTO project_node_fixed(project_id,std_node_id,fixed_date) VALUES(?,?,?)",
                        (pid, nid, fixed_date))
            log_rule_change("项目", nid, pid, "固定日期", before, {"fixed_date": fixed_date})
        elif m == "rule":
            try:
                branches = json.loads(request.form.get("rules_json", "[]"))
            except Exception:
                branches = []
            clean = sanitize_branches(branches)
            cur.execute("DELETE FROM project_node_fixed WHERE project_id=? AND std_node_id=?", (pid, nid))
            cur.execute("DELETE FROM project_node_rule WHERE project_id=? AND std_node_id=?", (pid, nid))
            for i, b in enumerate(clean):
                cur.execute(
                    "INSERT INTO project_node_rule(project_id,std_node_id,branch_order,is_default,conditions_json,depend_row,offset_expr,direction,base_is_t0) "
                    "VALUES(?,?,?,?,?,?,?,?,0)",
                    (pid, nid, i + 1, 1 if b["is_default"] else 0,
                     json.dumps(b["conditions"], ensure_ascii=False), b["depend_row"], b["offset_expr"], b["direction"]))
            log_rule_change("项目", nid, pid, "更新规则", before, clean_to_front(clean))
        else:
            con.close(); return jsonify(ok=False, msg="未知模式"), 400
        con.commit(); con.close()
        con = db(); cur = con.cursor()
        cur.execute("SELECT plan_start,prereq_json FROM project WHERE id=?", (pid,))
        pj = cur.fetchone(); con.close()
        # 单节点调整后重新生成三版（保留 actual_*/status，并刷新三版预警）
        regenerate_all_versions(pid)
        return jsonify(ok=True, mode=m)

    con = db(); cur = con.cursor()
    cur.execute("SELECT id,action,operator,created_at FROM rule_change_log "
                "WHERE scope='项目' AND project_id=? AND std_node_id=? ORDER BY id DESC LIMIT 20", (pid, nid))
    logs = [dict(r) for r in cur.fetchall()]; cur.close(); con.close()
    return jsonify(dict(nid=nid, name=node["name"], level=node["level"], seq=node["seq"],
                        has_rule=has_rule, current_mode=mode, fixed_date=fixed_date,
                        current_date=fin["finish_inner"] if fin else None, rules=rules, logs=logs))


# ---------------- 实际进度填报（独立页） ----------------
@app.route("/project/<int:pid>/fill", methods=["GET", "POST"])
def project_fill(pid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name,client,status FROM project WHERE id=?", (pid,))
    p = cur.fetchone()
    if not p:
        cur.close(); con.close(); return "项目不存在", 404
    if request.method == "POST":
        con.close()
        actual = []
        snids = set()
        for k in request.form:
            for pref in ("actual_start_", "actual_finish_", "progress_"):
                if k.startswith(pref):
                    snids.add(k[len(pref):])
        for snid in snids:
            snid_i = int(snid)
            as_ = (request.form.get(f"actual_start_{snid}") or "").strip()
            af = (request.form.get(f"actual_finish_{snid}") or "").strip()
            prog = (request.form.get(f"progress_{snid}") or "").strip()
            as_v = _as_date(as_) if as_ else None
            af_v = _as_date(af) if af else None
            prog_v = int(prog) if (prog.isdigit()) else 0
            prog_v = max(0, min(100, prog_v))
            actual.append(dict(std_node_id=snid_i,
                                actual_start=as_v.isoformat() if as_v else None,
                                actual_finish=af_v.isoformat() if af_v else None,
                                progress_pct=prog_v))
        save_actual(pid, actual)
        return redirect(url_for("project_detail", pid=pid))
    # GET：列出内控版全部节点，预填已有 actual_*
    ensure_all_versions(pid)
    vid = version_vid(pid, "内控")
    cur.execute(
        "SELECT pn.std_node_id,pn.level,pn.seq,pn.name,pn.finish_inner,pn.actual_start,pn.actual_finish,"
        "pn.progress_pct,pn.status,pr.name AS prof,sn.excel_row "
        "FROM plan_node pn LEFT JOIN profession pr ON pn.profession_id=pr.id "
        "LEFT JOIN std_node sn ON pn.std_node_id=sn.id "
        "WHERE pn.project_id=? AND pn.version_id=? ORDER BY sn.excel_row", (pid, vid))
    rows = [dict(r) for r in cur.fetchall()]
    con.close()
    return render_template("project_fill.html", p=p, rows=rows, today=TODAY)


# ---------------- 三版计划对比 ----------------
@app.route("/project/<int:pid>/compare")
def project_compare(pid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name,client FROM project WHERE id=?", (pid,))
    p = cur.fetchone()
    if not p:
        cur.close(); con.close(); return "项目不存在", 404
    ensure_all_versions(pid)
    rows = compare_rows(pid)
    con.close()
    return render_template("project_compare.html", p=p, rows=rows, today=TODAY)


# ---------------- 导出 Excel（三版计划对比） ----------------
@app.route("/project/<int:pid>/export.xlsx")
def project_export_xlsx(pid):
    buf = build_compare_xlsx(pid)
    if buf is None:
        return "项目不存在", 404
    con = db(); cur = con.cursor()
    cur.execute("SELECT name FROM project WHERE id=?", (pid,))
    p = cur.fetchone(); con.close()
    fname = f"{(p['name'] if p else 'project')}_三版计划对比_{TODAY}.xlsx"
    return send_file(buf,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=fname)


# ---------------- 预警引擎 ----------------
LEVELS = ["里程碑", "一级", "二级", "二级*"]
PLAN_TYPES = ["内控", "考核", "履约"]


@app.route("/alert-rules")
def alert_rules():
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,level,plan_type,threshold_days,notify_target FROM alert_rule ORDER BY plan_type,level")
    rules = [dict(r) for r in cur.fetchall()]
    cur.close(); con.close()
    return render_template("alert_rules.html", rules=rules, levels=LEVELS, plan_types=PLAN_TYPES)


@app.route("/alert-rule/new", methods=["GET", "POST"])
@app.route("/alert-rule/<int:rid>/edit", methods=["GET", "POST"])
def alert_rule_edit(rid=None):
    con = db(); cur = con.cursor()
    if request.method == "POST":
        level = request.form["level"]
        plan_type = request.form["plan_type"]
        threshold_days = int(request.form["threshold_days"])
        notify_target = (request.form.get("notify_target") or "").strip()
        if rid:
            cur.execute(
                "UPDATE alert_rule SET level=?,plan_type=?,threshold_days=?,notify_target=? WHERE id=?",
                (level, plan_type, threshold_days, notify_target, rid))
        else:
            cur.execute(
                "INSERT INTO alert_rule(level,plan_type,threshold_days,notify_target) VALUES(?,?,?,?)",
                (level, plan_type, threshold_days, notify_target))
        con.commit(); cur.close(); con.close()
        return redirect(url_for("alert_rules"))
    rule = None
    if rid:
        cur.execute("SELECT id,level,plan_type,threshold_days,notify_target FROM alert_rule WHERE id=?", (rid,))
        rule = cur.fetchone()
    cur.close(); con.close()
    return render_template("alert_rule_form.html", rule=rule, levels=LEVELS, plan_types=PLAN_TYPES)


@app.route("/alert-rule/<int:rid>/delete", methods=["POST"])
def alert_rule_delete(rid):
    con = db(); cur = con.cursor()
    cur.execute("DELETE FROM alert_record WHERE alert_rule_id=?", (rid,))
    cur.execute("DELETE FROM alert_rule WHERE id=?", (rid,))
    con.commit(); cur.close(); con.close()
    return redirect(url_for("alert_rules"))


@app.route("/project/<int:pid>/compute-alerts", methods=["POST"])
def project_compute_alerts(pid):
    plan_type = request.form.get("plan_type", "内控")
    if plan_type == "全部":
        for pt in PLAN_TYPES:
            compute_alerts(pid, pt, as_of=TODAY)
    else:
        compute_alerts(pid, plan_type, as_of=TODAY)
    return redirect(url_for("project_alerts", pid=pid))


@app.route("/project/<int:pid>/alerts")
def project_alerts(pid):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,name FROM project WHERE id=?", (pid,))
    p = cur.fetchone()
    if not p:
        cur.close(); con.close(); return "项目不存在", 404
    cur.execute(
        "SELECT ar.id,ar.plan_node_id,ar.triggered_at,ar.delay_days,ar.message,pn.level,pn.name AS node_name,"
        "pn.status AS node_status,pn.finish_inner "
        "FROM alert_record ar JOIN plan_node pn ON ar.plan_node_id=pn.id "
        "WHERE pn.project_id=? ORDER BY ar.delay_days DESC", (pid,))
    records = [dict(r) for r in cur.fetchall()]
    cur.execute("SELECT COUNT(*) AS c FROM plan_node WHERE project_id=? AND status<>'正常'", (pid,))
    flagged = cur.fetchone()["c"]
    cur.close(); con.close()
    return render_template("project_alerts.html", p=p, records=records, flagged=flagged,
                           plan_types=PLAN_TYPES, today=TODAY)


@app.route("/alerts")
def alert_center():
    con = db(); cur = con.cursor()
    cur.execute("SELECT id FROM project ORDER BY id")
    pids = [r["id"] for r in cur.fetchall()]
    for pid in pids:
        for pt in PLAN_TYPES:
            compute_alerts(pid, pt, as_of=TODAY)
    cur.execute(
        "SELECT ar.id,ar.plan_node_id,ar.delay_days,ar.message,ar.triggered_at,pn.level,pn.name AS node_name,"
        "pn.status AS node_status,pr.name AS project_name,pr.id AS project_id "
        "FROM alert_record ar JOIN plan_node pn ON ar.plan_node_id=pn.id "
        "JOIN project pr ON pn.project_id=pr.id ORDER BY ar.delay_days DESC")
    records = [dict(r) for r in cur.fetchall()]
    cur.close(); con.close()
    return render_template("alert_center.html", records=records, total=len(records),
                           plan_types=PLAN_TYPES, today=TODAY)


# ---------------- 标准库版本化：发布 / 回滚 / 查看 ----------------
def cur_work_version(cur):
    """当前工作版本号 = std_rule_version 最大 version_no + 1；无记录则 1。"""
    cur.execute("SELECT COALESCE(MAX(version_no),0) AS m FROM std_rule_version")
    return (cur.fetchone()["m"] or 0) + 1


def snapshot_active_rules():
    """读取当前生效集（is_active=1）全部规则到内存，供发布/回滚复制（避免同表自覆盖）。"""
    con = db(); cur = con.cursor()
    cur.execute("SELECT std_node_id,branch_order,is_default,conditions_json,depend_row,offset_expr,direction,base_is_t0 "
                "FROM std_node_rule WHERE is_active=1 ORDER BY std_node_id, branch_order")
    rules = [dict(r) for r in cur.fetchall()]
    cur.execute("SELECT std_node_id,depend_std_node_id,depend_row,offset_days,direction,base_is_t0 "
                "FROM std_node_dependency WHERE is_active=1")
    deps = [dict(r) for r in cur.fetchall()]
    cur.close(); con.close()
    return rules, deps


def publish_version(name, note, effective_from, operator="代建运营部"):
    """把当前工作副本固化为新版本快照（整库全量复制，旧版 is_active=0）。"""
    con = db(); cur = con.cursor()
    new_v = cur_work_version(cur)
    rules, deps = snapshot_active_rules()
    cur.execute("UPDATE std_node_rule SET is_active=0 WHERE is_active=1")
    cur.execute("UPDATE std_node_dependency SET is_active=0 WHERE is_active=1")
    for r in rules:
        cur.execute(
            "INSERT INTO std_node_rule(std_node_id,branch_order,is_default,conditions_json,depend_row,offset_expr,direction,base_is_t0,version_no,is_active) "
            "VALUES(?,?,?,?,?,?,?,?,?,1)",
            (r["std_node_id"], r["branch_order"], r["is_default"], r["conditions_json"], r["depend_row"],
             r["offset_expr"], r["direction"], r["base_is_t0"], new_v))
    for d in deps:
        cur.execute(
            "INSERT INTO std_node_dependency(std_node_id,depend_std_node_id,depend_row,offset_days,direction,base_is_t0,version_no,is_active) "
            "VALUES(?,?,?,?,?,?,?,1)",
            (d["std_node_id"], d["depend_std_node_id"], d["depend_row"], d["offset_days"], d["direction"],
             d["base_is_t0"], new_v))
    cur.execute(
        "INSERT INTO std_rule_version(version_no,name,note,effective_from,operator,change_summary) "
        "VALUES(?,?,?,?,?,?)",
        (new_v, name or f"版本{new_v}", note, effective_from, operator, "发布当前工作副本"))
    con.commit(); cur.close(); con.close()
    return new_v


def rollback_version(v, operator="代建运营部"):
    """回滚到版本 v：复制 v 的全部规则为最新版本并生效，写变更日志（传 con 防 SQLite 锁）。"""
    con = db(); cur = con.cursor()
    new_v = cur_work_version(cur)
    before = snapshot_active_rules()
    cur.execute("SELECT std_node_id,branch_order,is_default,conditions_json,depend_row,offset_expr,direction,base_is_t0 "
                "FROM std_node_rule WHERE version_no=?", (v,))
    rules = [dict(r) for r in cur.fetchall()]
    cur.execute("SELECT std_node_id,depend_std_node_id,depend_row,offset_days,direction,base_is_t0 "
                "FROM std_node_dependency WHERE version_no=?", (v,))
    deps = [dict(r) for r in cur.fetchall()]
    if not rules and not deps:
        cur.close(); con.close()
        raise ValueError(f"版本 v{v} 不存在")
    cur.execute("UPDATE std_node_rule SET is_active=0 WHERE is_active=1")
    cur.execute("UPDATE std_node_dependency SET is_active=0 WHERE is_active=1")
    for r in rules:
        cur.execute(
            "INSERT INTO std_node_rule(std_node_id,branch_order,is_default,conditions_json,depend_row,offset_expr,direction,base_is_t0,version_no,is_active) "
            "VALUES(?,?,?,?,?,?,?,?,?,1)",
            (r["std_node_id"], r["branch_order"], r["is_default"], r["conditions_json"], r["depend_row"],
             r["offset_expr"], r["direction"], r["base_is_t0"], new_v))
    for d in deps:
        cur.execute(
            "INSERT INTO std_node_dependency(std_node_id,depend_std_node_id,depend_row,offset_days,direction,base_is_t0,version_no,is_active) "
            "VALUES(?,?,?,?,?,?,?,1)",
            (d["std_node_id"], d["depend_std_node_id"], d["depend_row"], d["offset_days"], d["direction"],
             d["base_is_t0"], new_v))
    cur.execute(
        "INSERT INTO std_rule_version(version_no,name,note,operator,source_version_no,change_summary) "
        "VALUES(?,?,?,?,?,?)",
        (new_v, f"回滚自 v{v}", f"从 v{v} 回滚生成 v{new_v}", operator, v, f"回滚到 v{v}"))
    after = snapshot_active_rules()
    log_rule_change("标准库", 0, None, "版本回滚",
                    {"from_version": v, "rules": before[0], "deps": before[1]},
                    {"to_version": new_v, "rules": after[0], "deps": after[1]},
                    operator=operator, con=con)
    con.commit(); cur.close(); con.close()
    return new_v


@app.route("/std-rule/publish", methods=["GET", "POST"])
def std_rule_publish():
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        note = (request.form.get("note") or "").strip()
        ef = request.form.get("effective_from") or None
        publish_version(name, note, ef)
        return redirect(url_for("std_nodes"))
    return render_template("std_rule_publish.html")


@app.route("/std-rule/rollback/<int:v>", methods=["POST"])
def std_rule_rollback(v):
    rollback_version(v)
    return redirect(url_for("std_nodes"))


@app.route("/std-rule/version/<int:v>")
def std_rule_version_view(v):
    con = db(); cur = con.cursor()
    cur.execute("SELECT id,excel_row,level,seq,name,profession_id,tmpl_finish_formula,depend_row,offset_days,base_is_t0 FROM std_node ORDER BY excel_row")
    nodes = cur.fetchall()
    cur.execute("SELECT id,name FROM profession")
    prof = {r["id"]: r["name"] for r in cur.fetchall()}
    cur.execute("SELECT id,excel_row,level,seq,name FROM std_node")
    node_meta = {r["id"]: r for r in cur.fetchall()}
    cur.execute("SELECT std_node_id,depend_row,offset_days,base_is_t0,direction,depend_std_node_id FROM std_node_dependency WHERE version_no=?", (v,))
    deps = {(r["std_node_id"]): r for r in cur.fetchall()}
    cur.execute("SELECT DISTINCT std_node_id FROM std_node_rule WHERE version_no=? AND is_active=1", (v,))
    rule_ids = {r["std_node_id"] for r in cur.fetchall()}
    cur.execute("SELECT name,effective_from,created_at FROM std_rule_version WHERE version_no=?", (v,))
    vrow = cur.fetchone()
    cur.close(); con.close()
    vname = vrow["name"] if vrow else None
    veffective = vrow["effective_from"] if vrow else None
    vcreated = vrow["created_at"] if vrow else None
    for n in nodes:
        n["prof"] = prof.get(n["profession_id"], "—")
        dp = deps.get(n["id"])
        if dp:
            if dp["base_is_t0"]:
                n["dep_desc"] = "T0基准" + (f" +{dp['offset_days']}天" if dp["offset_days"] else "")
            else:
                dn = node_meta.get(dp["depend_std_node_id"])
                dname = dn["name"] if dn else f"L{dp['depend_row']}"
                n["dep_desc"] = (f"依赖 {dname}(L{dp['depend_row']}) "
                                 f"{dp['direction']} {abs(dp['offset_days'])}天")
        else:
            n["dep_desc"] = "—"
        if n["id"] in rule_ids:
            n["dep_desc"] += " · 含参数化分支"
    return render_template("std_rule_version.html", nodes=nodes, total=len(nodes),
                           v=v, vname=vname, veffective=veffective, vcreated=vcreated)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    # 注意：公网/预览环境务必 debug=False，避免 Werkzeug 调试器远程代码执行风险
    app.run(host="0.0.0.0", port=port, debug=False)
